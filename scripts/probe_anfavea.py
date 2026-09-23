"""TEMP diagnostic for the ANFAVEA site relaunch — removed before merge."""
import io
import os
import re

import openpyxl
import requests
from bs4 import BeautifulSoup

H = {
    "User-Agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                   "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.8",
}
OUT = "probe_out"
os.makedirs(OUT, exist_ok=True)
PAGES = [
    "https://anfavea.com.br/site/edicoes-em-excel/",
    "https://anfavea.com.br/site/issues-in-excel/",
    "https://anfavea.com.br/site/",
    "https://anfavea.com.br/",
]
cands = []
for i, url in enumerate(PAGES):
    try:
        r = requests.get(url, headers=H, timeout=30)
    except Exception as e:
        print("ERR", url, e)
        continue
    print(f"\n=== {url} -> {r.status_code} {r.url} len={len(r.text)}")
    open(f"{OUT}/page{i}.html", "w").write(r.text)
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup.find_all(True):
        for attr in ("href", "src", "data-href", "data-url", "data-src", "action"):
            v = tag.get(attr)
            if v and re.search(r"xls|csv|zip|excel|estat|emplac|download|drive|sharepoint|onedrive|dados|\.pdf|api|json", v, re.I):
                print(f"  <{tag.name} {attr}> {v}  | text={tag.get_text(' ', strip=True)[:60]!r}")
                if re.search(r"\.xlsx?(\?|$)", v, re.I):
                    cands.append(v if v.startswith("http") else "https://anfavea.com.br" + v)
    for m in set(re.findall(r"https?://[^\s\"'<>]+?\.xlsx?", r.text, re.I)):
        print("  RAW xls:", m)
        cands.append(m)
    # visible text around "Excel"/"2026"
    txt = soup.get_text("\n", strip=True)
    for line in txt.splitlines():
        if re.search(r"2026|2025|excel|planilha|séries|series", line, re.I):
            print("   txt:", line[:120])

seen = set()
for u in cands:
    if u in seen or not re.search(r"202[4-6]", u):
        continue
    seen.add(u)
    print(f"\n##### {u}")
    try:
        r = requests.get(u, headers=H, timeout=60)
        print("status", r.status_code, "bytes", len(r.content), r.headers.get("content-type"))
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True, read_only=True)
    except Exception as e:
        print("  ERR", e)
        continue
    for ws in wb.worksheets:
        print(f"--- sheet {ws.title!r} dims={ws.max_row}x{ws.max_column}")
        for j, row in enumerate(ws.iter_rows(values_only=True)):
            if j >= 45:
                break
            cells = [str(c)[:18] for c in row[:16]]
            while cells and cells[-1] == "None":
                cells.pop()
            if cells:
                print(f"  {j:3d}|", " | ".join(cells))
